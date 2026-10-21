"""Chuyển Markdown đơn giản sang .docx để nộp báo cáo tuần.

Hỗ trợ: tiêu đề #..####, đoạn văn, danh sách -, *, 1., bảng |..|, khối ```code```,
ảnh ![chú thích](đường_dẫn), **đậm**, *nghiêng*, `code`.

Dùng: python md2docx.py input.md [output.docx]
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BLACK = RGBColor(0, 0, 0)  # mọi chữ trong báo cáo dùng một màu đen

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)")


def add_inline(par, text, size=None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = par.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = par.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = BLACK
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = par.add_run(part[1:-1])
            run.italic = True
        else:
            run = par.add_run(part)
        if size:
            run.font.size = Pt(size)


def shade(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def add_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    header, body = cells[0], [r for r in cells[1:] if not set("".join(r)) <= set("-: ")]
    ncol = len(header)
    table = doc.add_table(rows=1 + len(body), cols=ncol)
    table.style = "Table Grid"
    font = 9 if ncol >= 5 else 10
    for j, h in enumerate(header):
        c = table.rows[0].cells[j]
        c.text = ""
        add_inline(c.paragraphs[0], h, font)
        for r in c.paragraphs[0].runs:
            r.bold = True
        shade(c, "E7E6E6")  # nền xám nhạt trung tính, chữ vẫn đen
    for i, r in enumerate(body, start=1):
        for j in range(ncol):
            c = table.rows[i].cells[j]
            c.text = ""
            add_inline(c.paragraphs[0], r[j] if j < len(r) else "", font)
    doc.add_paragraph()


def convert(md_path, out_path):
    md_path = Path(md_path)
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = BLACK
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    for name in ("Heading 1", "Heading 2", "Heading 3", "Heading 4", "Title"):
        st = doc.styles[name]
        st.font.name = "Times New Roman"
        rfonts = st.element.rPr.rFonts
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rfonts.attrib.pop(qn(attr), None)  # bỏ font theme (Calibri Light) để dùng Times New Roman
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(attr), "Times New Roman")
        st.font.color.rgb = BLACK

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if s.startswith("```"):
            buf = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            run = p.add_run("\n".join(buf))
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            i += 1
            continue
        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i])
                i += 1
            add_table(doc, rows)
            continue
        m = re.match(r"!\[(.*?)\]\((.*?)\)", s)
        if m:
            img = (md_path.parent / m.group(2)).resolve()
            if img.exists():
                doc.add_picture(str(img), width=Cm(16))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = cap.add_run(m.group(1))
                r.italic = True
                r.font.size = Pt(10)
            i += 1
            continue
        h = re.match(r"^(#{1,4})\s+(.*)", s)
        if h:
            level = len(h.group(1))
            if level == 1 and i == 0:
                t = doc.add_heading(h.group(2), level=0)
                t.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                doc.add_heading(h.group(2), level=min(level, 3) if level > 1 else 1)
            i += 1
            continue
        b = re.match(r"^(\s*)[-*]\s+(.*)", line)
        if b and not s.startswith("**") or (b and s.startswith("* ")):
            depth = len(b.group(1)) // 2
            p = doc.add_paragraph(style="List Bullet" if depth == 0 else "List Bullet 2")
            add_inline(p, b.group(2))
            i += 1
            continue
        n = re.match(r"^\s*\d+[.)]\s+(.*)", line)
        if n:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, n.group(1))
            i += 1
            continue
        if s.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            add_inline(p, s.lstrip("> "))
            for r in p.runs:
                r.italic = True
            i += 1
            continue
        if not s:
            i += 1
            continue
        buf = [s]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^\s*(#|[-*]\s|\d+[.)]\s|\||```|!\[|>)", lines[i]):
            buf.append(lines[i].strip())
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        add_inline(p, " ".join(buf))
    force_black(doc)
    doc.save(out_path)
    return out_path


def force_black(doc):
    """Ép mọi màu chữ (cả trong styles mặc định của mẫu Word) về đen, bỏ màu theo theme."""
    for root in (doc.styles.element, doc.element):
        for el in root.iter(qn("w:color")):
            for attr in ("w:themeColor", "w:themeShade", "w:themeTint"):
                el.attrib.pop(qn(attr), None)
            el.set(qn("w:val"), "000000")


if __name__ == "__main__":
    src = sys.argv[1]
    dst = sys.argv[2] if len(sys.argv) > 2 else str(Path(src).with_suffix(".docx"))
    print(convert(src, dst))
